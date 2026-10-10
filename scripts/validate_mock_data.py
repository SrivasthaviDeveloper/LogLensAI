"""Validate the synthetic workbook, registry, splits, manifest, and JSONL traces."""

from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.mock_dataset import (  # noqa: E402
    DEFAULT_OUTPUT, DICTIONARY_COLUMNS, REGISTRY_COLUMNS, SCENARIO_COLUMNS, TRAINING_COLUMNS,
)


def _rows(sheet) -> tuple[list[str], list[dict[str, Any]]]:
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return [], []
    headers = [str(value) for value in values[0]]
    rows = [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)]
    return headers, rows


def validate_dataset(output_root: Path = DEFAULT_OUTPUT) -> list[str]:
    """Return validation issues; an empty list means all synthetic invariants hold."""
    root = Path(output_root).resolve()
    errors: list[str] = []
    workbook_path = root / "loglens_training_data.xlsx"
    manifest_path = root / "upload_manifest.json"
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except Exception as exc:
        return [f"Training workbook could not be opened: {type(exc).__name__}."]
    expected_sheets = ["TrainingData", "ServiceRegistry", "ScenarioCatalog", "DataDictionary"]
    if workbook.sheetnames != expected_sheets:
        errors.append("Workbook sheets do not match the required sheet names and order.")
    training_headers, training = _rows(workbook["TrainingData"])
    registry_headers, registry_rows = _rows(workbook["ServiceRegistry"])
    scenario_headers, scenarios = _rows(workbook["ScenarioCatalog"])
    dictionary_headers, dictionary = _rows(workbook["DataDictionary"])
    for sheet_name, actual, required in (
        ("TrainingData", training_headers, TRAINING_COLUMNS),
        ("ServiceRegistry", registry_headers, REGISTRY_COLUMNS),
        ("ScenarioCatalog", scenario_headers, SCENARIO_COLUMNS),
        ("DataDictionary", dictionary_headers, DICTIONARY_COLUMNS),
    ):
        if not set(required).issubset(actual):
            errors.append(f"{sheet_name} is missing required columns: {sorted(set(required) - set(actual))}.")
    workbook.close()

    if len(scenarios) < 60:
        errors.append("Expected at least 20 scenarios per each of the three applications.")
    scenario_ids = [str(row.get("scenario_id") or "") for row in scenarios]
    if len(scenario_ids) != len(set(scenario_ids)):
        errors.append("Scenario IDs are not unique.")
    record_ids = [str(row.get("record_id") or "") for row in training]
    if len(record_ids) != len(set(record_ids)):
        errors.append("Training record IDs are not unique.")
    request_ids = [str(row["expected_request_id"]) for row in scenarios if row.get("expected_request_id")]
    if len(request_ids) != len(set(request_ids)):
        errors.append("Expected request IDs are not unique.")
    if any(row.get("is_synthetic") is not True for row in training):
        errors.append("Every training record must be marked is_synthetic=TRUE.")
    training_scenarios = {str(row.get("scenario_id")) for row in training}
    if training_scenarios != set(scenario_ids):
        errors.append("TrainingData and ScenarioCatalog scenario IDs differ.")

    split_by_scenario: dict[str, set[str]] = defaultdict(set)
    for row in training:
        split_by_scenario[str(row.get("scenario_id"))].add(str(row.get("data_split")))
    for row in scenarios:
        split_by_scenario[str(row.get("scenario_id"))].add(str(row.get("data_split")))
    if any(len(splits) != 1 for splits in split_by_scenario.values()):
        errors.append("A scenario group leaks across training, validation, or test splits.")
    if not any("test" in splits for splits in split_by_scenario.values()):
        errors.append("No held-out test scenarios were generated.")

    service_index = {
        (str(row.get("application_name")), str(row.get("service_name"))): row
        for row in registry_rows
    }
    if len(service_index) != len(registry_rows):
        errors.append("Service registry contains duplicate application/service keys.")
    dictionary_fields = {str(row.get("field_name")) for row in dictionary}
    workbook_data_fields = set(training_headers + registry_headers + scenario_headers)
    if not workbook_data_fields.issubset(dictionary_fields):
        errors.append(f"DataDictionary is missing fields: {sorted(workbook_data_fields - dictionary_fields)}.")

    manifest: dict[str, Any] = {}
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        errors.append("Upload manifest is missing or invalid JSON.")
    file_by_service = {
        (str(row["application_name"]), str(row["service_name"])): f"{row['log_directory']}/app.log"
        for row in registry_rows
    }
    events_by_scenario: dict[str, list[dict[str, Any]]] = defaultdict(list)
    seen_scenario_log_files: set[str] = set()
    for file_entry in manifest.get("files", []):
        relative = str(file_entry.get("relative_path", ""))
        file_path = (root / "mock_logs" / Path(relative)).resolve()
        if not file_path.is_relative_to((root / "mock_logs").resolve()) or not file_path.is_file():
            errors.append(f"Manifest log path is missing or outside mock_logs: {relative}.")
            continue
        data = file_path.read_bytes()
        import hashlib
        if hashlib.sha256(data).hexdigest() != file_entry.get("sha256"):
            errors.append(f"Manifest checksum does not match {relative}.")
        line_count = 0
        for line_number, line in enumerate(data.decode("utf-8").splitlines(), start=1):
            if not line.strip():
                continue
            line_count += 1
            try:
                event = json.loads(line)
            except json.JSONDecodeError:
                errors.append(f"Invalid JSONL record at {relative}:{line_number}.")
                continue
            try:
                stamp = datetime.fromisoformat(event["timestamp"].replace("Z", "+00:00"))
                if stamp.tzinfo is None or stamp.utcoffset() != timezone.utc.utcoffset(stamp):
                    errors.append(f"Timestamp is not timezone-aware UTC at {relative}:{line_number}.")
            except (KeyError, TypeError, ValueError):
                errors.append(f"Timestamp is missing or invalid at {relative}:{line_number}.")
            if event.get("is_synthetic") is not True:
                errors.append(f"Log record is not marked synthetic at {relative}:{line_number}.")
            app, service = str(event.get("application_name")), str(event.get("service_name"))
            registry_row = service_index.get((app, service))
            if registry_row is None:
                errors.append(f"Log service {app}/{service} is absent from ServiceRegistry.")
            elif relative != f"{registry_row['log_directory']}/app.log":
                errors.append(f"Log path does not match ServiceRegistry for {app}/{service}.")
            scenario_id = str(event.get("scenario_id") or "")
            events_by_scenario[scenario_id].append(event)
        if line_count != int(file_entry.get("event_count", -1)):
            errors.append(f"Manifest event count does not match {relative}.")

    catalog_by_id = {str(row.get("scenario_id")): row for row in scenarios}
    if set(events_by_scenario) != set(catalog_by_id):
        errors.append("Generated JSONL scenario IDs do not exactly match ScenarioCatalog.")
    for scenario_id, scenario in catalog_by_id.items():
        events = events_by_scenario.get(scenario_id, [])
        if len(events) < 2:
            errors.append(f"Scenario {scenario_id} does not have multiple log events.")
            continue
        if not any(event.get("event_type") == "request" for event in events):
            errors.append(f"Scenario {scenario_id} has no request event.")
        expected_request_id = scenario.get("expected_request_id")
        expected_id = str(expected_request_id) if expected_request_id else None
        if any(event.get("request_id") != expected_id or event.get("correlation_id") != expected_id for event in events):
            errors.append(f"Request/correlation IDs are inconsistent for {scenario_id}.")
        if not any(event.get("service_name") == scenario.get("expected_service") for event in events):
            errors.append(f"Expected service has no event for {scenario_id}.")
        expected_status = scenario.get("expected_http_status")
        responses = [event for event in events if event.get("event_type") == "response"]
        if expected_status is None:
            if responses:
                errors.append(f"Incomplete scenario {scenario_id} unexpectedly contains a response event.")
        elif not any(event.get("http_status") == expected_status for event in responses):
            errors.append(f"Expected HTTP status is absent from response events for {scenario_id}.")

    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    issues = validate_dataset(args.output_root)
    if issues:
        print("Synthetic data validation failed:")
        for issue in issues:
            print(f"- {issue}")
        return 1
    manifest = json.loads((args.output_root / "upload_manifest.json").read_text(encoding="utf-8"))
    print(f"Valid: {manifest['scenario_count']} scenarios, {manifest['event_count']} log events, {len(manifest['files'])} files.")
    print("Every record is synthetic; scenario IDs remain grouped by train/validation/test split.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
