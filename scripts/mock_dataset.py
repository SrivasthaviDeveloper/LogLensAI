"""Shared fictional scenario and workbook generation helpers."""

from __future__ import annotations

import hashlib
import json
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "data" / "training"
REGISTRY_PATH = ROOT / "config" / "training_service_registry.json"
APP_DOMAINS = {
    "VNSP": "VDC (Virtual Data Center) details.",
    "ESP": "entity details.",
    "IPSM": "customer IP details.",
}
TRAINING_COLUMNS = [
    "record_id", "application_name", "application_domain", "error_message", "error_code",
    "environment", "service_name", "server_id", "request_pattern", "response_pattern",
    "http_status", "historical_root_cause", "resolution", "scenario_id", "is_synthetic",
    "data_split",
]
REGISTRY_COLUMNS = [
    "application_name", "service_name", "server_id", "environment", "log_directory",
    "enabled", "service_dependencies", "description",
]
SCENARIO_COLUMNS = [
    "scenario_id", "application_name", "scenario_name", "input_error", "expected_service",
    "expected_error_category", "expected_root_cause", "expected_request_id",
    "expected_http_status", "expected_evidence", "expected_resolution", "expected_result_status", "data_split",
]
DICTIONARY_COLUMNS = ["field_name", "meaning", "expected_type", "role", "notes"]

# Tuple order: name, error wording, code, category, service, status, mode, cause, resolution.
SCENARIO_TEMPLATES: dict[str, list[tuple[str, str, str, str, str, int | None, str, str, str]]] = {
    "VNSP": [
        ("vdc_details_success", "VDC details were returned successfully", "", "success", "vdc-service", 200, "success", "No failure was observed in this synthetic request.", "No remediation is indicated by the successful fixture."),
        ("vdc_details_success_alt", "The requested virtual data center detail lookup completed", "", "success", "vdc-service", 200, "success", "No failure was observed in this synthetic request.", "No remediation is indicated by the successful fixture."),
        ("vdc_lookup_not_found", "VDC identifier was not found in the demo dataset", "VDC_NOT_FOUND", "not_found", "vdc-service", 404, "service_error", "The fictional VDC identifier has no matching demo record.", "Check the supplied demo identifier and fixture catalog."),
        ("vdc_record_missing", "No virtual data center details matched the sample reference", "VDC_NOT_FOUND", "not_found", "vdc-service", 404, "service_error", "The synthetic lookup returned no matching VDC record.", "Confirm the demo reference and expected fixture contents."),
        ("vdc_identifier_invalid", "VDC identifier failed the mock format validation", "VDC_ID_INVALID", "validation_error", "vdc-service", 400, "service_error", "The fictional identifier does not match the demo input format.", "Review the documented synthetic identifier format."),
        ("vdc_identifier_blank", "VDC reference is missing from the synthetic request", "VDC_REFERENCE_REQUIRED", "validation_error", "api-gateway", 400, "gateway_error", "The sample request omitted its required fictional VDC reference.", "Supply a non-empty demo identifier."),
        ("vdc_authentication_rejected", "Demo caller authentication was rejected", "AUTH_INVALID_DEMO_TOKEN", "authentication_error", "auth-service", 401, "auth_error", "The synthetic authorization fixture rejects the demo caller.", "Use the valid fictional authorization fixture; do not enter real credentials."),
        ("vdc_scope_forbidden", "Demo caller does not have the requested read scope", "AUTH_SCOPE_MISSING", "authorization_error", "auth-service", 403, "auth_error", "The synthetic caller lacks the fixture's requested scope.", "Compare the fictional required and granted scopes."),
        ("vdc_dependency_unavailable", "Mock dependency did not respond to the VDC lookup", "VDC_DEPENDENCY_UNAVAILABLE", "dependency_unavailable", "mock-dependency", 503, "dependency_error", "The synthetic dependency fixture is marked unavailable.", "Inspect correlated dependency and caller events in this demo trace."),
        ("vdc_upstream_bad_gateway", "Gateway could not establish a mock downstream response", "VDC_UPSTREAM_FAILURE", "gateway_error", "api-gateway", 502, "dependency_error", "The gateway fixture did not receive a valid downstream response.", "Compare the gateway and downstream events for the same demo request."),
        ("vdc_timeout_initial", "VDC details lookup exceeded the synthetic time budget", "VDC_LOOKUP_TIMEOUT", "request_timeout", "vdc-service", 504, "timeout", "The fictional VDC service exceeded the demo timeout threshold.", "Compare event timestamps and the configured synthetic timeout."),
        ("vdc_connection_refused", "Connection to the mock VDC service was refused", "VDC_CONNECTION_REFUSED", "connection_failure", "api-gateway", 502, "connection_refused", "The simulated gateway connection was refused before service receipt.", "Check the simulated service-availability state and trace boundary."),
        ("vdc_service_exception", "Synthetic VDC detail processing raised an internal error", "VDC_INTERNAL_ERROR", "internal_error", "vdc-service", 500, "service_error", "The demo service emitted a controlled internal-error event.", "Inspect the first correlated service error before proposing remediation."),
        ("vdc_not_found_alt", "Sample VDC reference has no configured mock record", "VDC_NOT_FOUND", "not_found", "vdc-service", 404, "service_error", "The synthetic fixture contains no matching VDC record.", "Check the demo reference against available fictional fixtures."),
        ("vdc_dependency_retry_exhausted", "VDC dependency remained unavailable after a mock retry", "VDC_RETRY_EXHAUSTED", "dependency_unavailable", "mock-dependency", 503, "dependency_error", "The deterministic dependency fixture exhausted its retry count.", "Review the recorded retry and final dependency event."),
        ("vdc_timeout_caller", "Caller deadline elapsed before VDC response arrived", "CALLER_DEADLINE_EXCEEDED", "request_timeout", "api-gateway", 504, "timeout", "The synthetic gateway deadline expired before a response event.", "Compare caller and service timestamps and timeout budgets."),
        ("vdc_payload_validation", "VDC query parameters failed mock validation", "VDC_QUERY_INVALID", "validation_error", "api-gateway", 400, "gateway_error", "The fictional request parameters violate the demo validation rule.", "Review only the supplied synthetic query fields."),
        ("vdc_processing_error", "VDC service returned a controlled processing error", "VDC_PROCESSING_ERROR", "internal_error", "vdc-service", 500, "service_error", "A deterministic demo exception was emitted by the VDC service.", "Inspect the correlated service event and response status."),
        ("vdc_trace_incomplete", "VDC request was logged without a response event", "RESPONSE_NOT_OBSERVED", "incomplete_evidence", "vdc-service", None, "incomplete", "The synthetic trace intentionally omits the response record.", "Search adjacent services for a matching request ID and time range."),
        ("vdc_correlation_missing", "VDC error sample has no reliable correlation identifier", "CORRELATION_ID_MISSING", "insufficient_evidence", "api-gateway", None, "insufficient", "The fixture omits request and correlation IDs, preventing a confirmed join.", "Provide a request ID or a narrower timestamp and additional logs."),
        ("vdc_timeout_second_wording", "Virtual data center detail response arrived after the allowed wait", "VDC_RESPONSE_LATE", "request_timeout", "vdc-service", 504, "timeout", "The fictional downstream response exceeded the allowed wait.", "Compare event timing and configured mock deadlines."),
        ("vdc_dependency_502", "Mock VDC dependency closed the connection unexpectedly", "VDC_DEPENDENCY_RESET", "dependency_unavailable", "mock-dependency", 502, "dependency_error", "The synthetic dependency fixture reset the downstream connection.", "Review dependency state and the matching gateway response."),
        ("vdc_success_repeat", "VDC detail request completed with a demo response", "", "success", "vdc-service", 200, "success", "No failure was observed in this synthetic request.", "No remediation is indicated by the successful fixture."),
        ("vdc_evidence_sparse", "VDC sample contains only an uncorrelated request event", "EVIDENCE_INCOMPLETE", "insufficient_evidence", "api-gateway", None, "insufficient", "The synthetic sample has no response or reliable request correlation.", "Provide a request ID and the downstream service log window."),
    ],
    "ESP": [
        ("entity_lookup_success", "Entity details were returned successfully", "", "success", "entity-service", 200, "success", "No failure was observed in this synthetic request.", "No remediation is indicated by the successful fixture."),
        ("entity_lookup_success_alt", "The sample entity-detail query completed", "", "success", "entity-service", 200, "success", "No failure was observed in this synthetic request.", "No remediation is indicated by the successful fixture."),
        ("entity_not_found", "Entity reference was not present in the fictional catalog", "ENTITY_NOT_FOUND", "not_found", "entity-service", 404, "service_error", "The synthetic entity identifier has no matching demo record.", "Check the supplied fictional entity reference."),
        ("entity_missing", "No entity details matched the demo lookup key", "ENTITY_NOT_FOUND", "not_found", "entity-service", 404, "service_error", "The demo entity lookup returned no matching record.", "Confirm the fixture identifier and expected sample data."),
        ("entity_validation_invalid", "Entity lookup key failed synthetic validation", "ENTITY_KEY_INVALID", "validation_error", "entity-service", 400, "service_error", "The fictional lookup key does not match the demo format.", "Review the documented synthetic entity-key format."),
        ("entity_validation_required", "Required entity reference was absent from the mock request", "ENTITY_REFERENCE_REQUIRED", "validation_error", "api-gateway", 400, "gateway_error", "The synthetic request omitted a required entity reference.", "Supply a non-empty fictional entity identifier."),
        ("entity_authentication_rejected", "Synthetic caller authentication was rejected", "AUTH_INVALID_DEMO_TOKEN", "authentication_error", "auth-service", 401, "auth_error", "The authorization fixture rejected the fictional caller.", "Use only the valid demo authorization fixture."),
        ("entity_scope_forbidden", "Synthetic caller lacks the entity-read permission", "AUTH_SCOPE_MISSING", "authorization_error", "auth-service", 403, "auth_error", "The sample caller does not have the fixture's requested scope.", "Compare the fictional required and granted scopes."),
        ("entity_dependency_unavailable", "Mock entity dependency was not available", "ENTITY_DEPENDENCY_UNAVAILABLE", "dependency_unavailable", "mock-dependency", 503, "dependency_error", "The synthetic dependency is marked unavailable for this trace.", "Inspect correlated dependency and caller events."),
        ("entity_upstream_bad_gateway", "Gateway did not receive a valid mock entity response", "ENTITY_UPSTREAM_FAILURE", "gateway_error", "api-gateway", 502, "dependency_error", "The gateway fixture did not receive a valid downstream response.", "Compare gateway and downstream events using the demo request ID."),
        ("entity_lookup_timeout", "Entity details lookup exceeded the synthetic deadline", "ENTITY_LOOKUP_TIMEOUT", "request_timeout", "entity-service", 504, "timeout", "The fictional entity service exceeded the configured demo deadline.", "Compare timestamps and the synthetic timeout configuration."),
        ("entity_connection_refused", "Connection to the mock entity service was refused", "ENTITY_CONNECTION_REFUSED", "connection_failure", "api-gateway", 502, "connection_refused", "The simulated gateway connection was refused before entity-service receipt.", "Check simulated service state and the request trace boundary."),
        ("entity_processing_exception", "Synthetic entity processing emitted an internal error", "ENTITY_INTERNAL_ERROR", "internal_error", "entity-service", 500, "service_error", "The demo service emitted a controlled exception event.", "Inspect the earliest correlated service error."),
        ("entity_not_found_alt", "Sample entity ID has no mock record", "ENTITY_NOT_FOUND", "not_found", "entity-service", 404, "service_error", "The synthetic entity fixture has no row for this key.", "Check the fictional key against the generated fixture set."),
        ("entity_trace_incomplete", "Entity request has no recorded response", "RESPONSE_NOT_OBSERVED", "incomplete_evidence", "entity-service", None, "incomplete", "The response event is intentionally missing from this demo trace.", "Search related service logs using the request ID and time range."),
        ("entity_correlation_missing", "Entity error sample lacks a reliable correlation ID", "CORRELATION_ID_MISSING", "insufficient_evidence", "api-gateway", None, "insufficient", "Request and correlation IDs are intentionally absent.", "Provide an exact request ID or additional timestamped log evidence."),
        ("entity_validation_alt", "Entity query contains an invalid demo field", "ENTITY_QUERY_INVALID", "validation_error", "api-gateway", 400, "gateway_error", "A fictional request field fails the mock validation rule.", "Review the synthetic request fields and validation event."),
        ("entity_dependency_retry", "Entity dependency remained unavailable after retry", "ENTITY_RETRY_EXHAUSTED", "dependency_unavailable", "mock-dependency", 503, "dependency_error", "The deterministic dependency fixture exhausted its configured retry.", "Review the retry event and terminal dependency status."),
        ("entity_timeout_caller", "Caller timed out while entity response was pending", "CALLER_DEADLINE_EXCEEDED", "request_timeout", "api-gateway", 504, "timeout", "The synthetic caller deadline expired before response completion.", "Compare caller and service timing in the correlated events."),
        ("entity_service_error", "Entity service returned a controlled server error", "ENTITY_PROCESSING_ERROR", "internal_error", "entity-service", 500, "service_error", "A deterministic internal-error event was generated for demonstration.", "Inspect the matching error and response events."),
        ("entity_auth_expired", "Fictional demo authorization was expired", "AUTH_DEMO_EXPIRED", "authentication_error", "auth-service", 401, "auth_error", "The test fixture marks its demo authorization as expired.", "Use the designated demo authorization fixture; never provide real secrets."),
        ("entity_scope_denied", "Entity lookup was denied by the synthetic scope check", "AUTH_SCOPE_DENIED", "authorization_error", "auth-service", 403, "auth_error", "The simulated authorization rule denies this sample operation.", "Compare the required demo scope with the configured sample scope."),
        ("entity_success_repeat", "Entity detail request completed with a demo response", "", "success", "entity-service", 200, "success", "No failure was observed in this synthetic request.", "No remediation is indicated by the successful fixture."),
        ("entity_evidence_sparse", "Entity sample has a request event but no correlated outcome", "EVIDENCE_INCOMPLETE", "insufficient_evidence", "api-gateway", None, "insufficient", "The synthetic sample lacks a reliable identifier and response.", "Provide a request ID and the relevant downstream log window."),
    ],
    "IPSM": [
        ("ip_lookup_success", "Fictional customer IP details were returned successfully", "", "success", "ip-service", 200, "success", "No failure was observed in this synthetic request.", "No remediation is indicated by the successful fixture."),
        ("ip_lookup_success_alt", "Sample customer IP record lookup completed", "", "success", "ip-service", 200, "success", "No failure was observed in this synthetic request.", "No remediation is indicated by the successful fixture."),
        ("ip_record_not_found", "IP record was not found in the fictional dataset", "IP_RECORD_NOT_FOUND", "not_found", "ip-service", 404, "service_error", "The reserved documentation address has no synthetic record.", "Check the fictional record reference and generated fixtures."),
        ("ip_customer_record_missing", "No customer IP detail matched the sample record key", "IP_RECORD_NOT_FOUND", "not_found", "ip-service", 404, "service_error", "The demo IP record lookup returned no matching fixture.", "Verify the fictional record key and fixture contents."),
        ("ip_input_invalid", "IP address input failed synthetic format validation", "IP_ADDRESS_INVALID", "validation_error", "ip-service", 400, "service_error", "The supplied synthetic input is not a valid IP address string.", "Validate the input format using documentation-range demo values."),
        ("ip_ipv4_format_invalid", "Mock IP lookup received a malformed address value", "IP_ADDRESS_FORMAT_INVALID", "validation_error", "api-gateway", 400, "gateway_error", "The mock request contains an invalid address representation.", "Review the synthetic address format before lookup."),
        ("ip_assignment_conflict", "Fictional IP address assignment conflicts with a demo record", "IP_ASSIGNMENT_CONFLICT", "assignment_conflict", "ip-service", 409, "service_error", "The synthetic fixture marks this documentation address as already assigned.", "Inspect the demo assignment record; do not change real customer data."),
        ("ip_duplicate_assignment", "Duplicate IP assignment detected in synthetic fixture", "IP_ASSIGNMENT_DUPLICATE", "assignment_conflict", "ip-service", 409, "service_error", "Two demo records reference the same reserved address.", "Review the generated synthetic assignment references."),
        ("ip_authentication_rejected", "Fictional caller authentication failed", "AUTH_INVALID_DEMO_TOKEN", "authentication_error", "auth-service", 401, "auth_error", "The simulated authorization fixture rejected the demo caller.", "Use only the valid fictional authorization fixture."),
        ("ip_read_forbidden", "Synthetic caller lacks permission to view IP details", "AUTH_SCOPE_MISSING", "authorization_error", "auth-service", 403, "auth_error", "The sample caller lacks the fixture's IP-read scope.", "Compare the fictional required and granted scopes."),
        ("ip_dependency_unavailable", "Mock address dependency is unavailable", "IP_DEPENDENCY_UNAVAILABLE", "dependency_unavailable", "mock-dependency", 503, "dependency_error", "The synthetic dependency fixture is marked unavailable.", "Inspect the dependency event and related caller events."),
        ("ip_upstream_bad_gateway", "Gateway received no valid mock IP-service response", "IP_UPSTREAM_FAILURE", "gateway_error", "api-gateway", 502, "dependency_error", "The synthetic gateway could not use the downstream response.", "Compare the correlated gateway and service events."),
        ("ip_lookup_timeout", "Customer IP detail lookup exceeded the demo timeout", "IP_LOOKUP_TIMEOUT", "request_timeout", "ip-service", 504, "timeout", "The fictional IP service exceeded its configured demo wait.", "Compare service duration with the synthetic deadline."),
        ("ip_connection_refused", "Connection to the mock IP service was refused", "IP_CONNECTION_REFUSED", "connection_failure", "api-gateway", 502, "connection_refused", "The simulated connection was refused before ip-service receipt.", "Check simulated endpoint state and trace evidence."),
        ("ip_internal_error", "Synthetic IP detail processing emitted a server error", "IP_INTERNAL_ERROR", "internal_error", "ip-service", 500, "service_error", "A controlled demo error was emitted by the fictional IP service.", "Inspect the first correlated service error and response."),
        ("ip_trace_incomplete", "IP detail request has no matching response event", "RESPONSE_NOT_OBSERVED", "incomplete_evidence", "ip-service", None, "incomplete", "The response record is intentionally omitted from this fixture.", "Search adjacent service logs by request ID and time range."),
        ("ip_correlation_missing", "IP error sample contains no reliable correlation identifier", "CORRELATION_ID_MISSING", "insufficient_evidence", "api-gateway", None, "insufficient", "The generated events intentionally omit request and correlation IDs.", "Provide a request ID or additional timestamped log evidence."),
        ("ip_not_found_alt", "Sample IP record reference has no mock entry", "IP_RECORD_NOT_FOUND", "not_found", "ip-service", 404, "service_error", "The synthetic dataset contains no matching IP record.", "Check the fictional reference against the generated sample."),
        ("ip_validation_alt", "IP detail query contains a malformed demo address", "IP_ADDRESS_INVALID", "validation_error", "api-gateway", 400, "gateway_error", "The mock request address fails the synthetic validation rule.", "Use a documentation-range IPv4 or IPv6 sample in this demo."),
        ("ip_assignment_conflict_alt", "Address assignment rejected because the mock record is occupied", "IP_ASSIGNMENT_CONFLICT", "assignment_conflict", "ip-service", 409, "service_error", "The controlled fixture models an address already assigned.", "Review the synthetic assignment trace for duplicate references."),
        ("ip_dependency_retry", "Address dependency remained unavailable after retry", "IP_RETRY_EXHAUSTED", "dependency_unavailable", "mock-dependency", 503, "dependency_error", "The deterministic dependency scenario exhausted its retry.", "Review the retry event and final dependency status."),
        ("ip_timeout_caller", "Caller deadline expired before IP details arrived", "CALLER_DEADLINE_EXCEEDED", "request_timeout", "api-gateway", 504, "timeout", "The mock gateway deadline expired while waiting for the service.", "Compare gateway and service durations in the same trace."),
        ("ip_lookup_success_repeat", "Customer IP demo detail request completed successfully", "", "success", "ip-service", 200, "success", "No failure was observed in this synthetic request.", "No remediation is indicated by the successful fixture."),
        ("ip_evidence_sparse", "IP sample contains no correlated response evidence", "EVIDENCE_INCOMPLETE", "insufficient_evidence", "api-gateway", None, "insufficient", "The generated request has no usable correlation or response event.", "Provide a request ID and matching service logs."),
    ],
}

DICTIONARY: list[tuple[str, str, str, str, str]] = [
    ("record_id", "Unique row identifier for one synthetic training example.", "string", "identifier", "Unique across the workbook."),
    ("application_name", "Fictional application label: VNSP, ESP, or IPSM.", "string", "feature", "Application meanings are defined by the project request."),
    ("application_domain", "Application meaning used to contextualize the example.", "string", "feature", "Fictional demo scope; not a production API description."),
    ("error_message", "Synthetic user-facing or log-style error wording.", "string", "feature", "All wording is invented for the simulation."),
    ("error_code", "Synthetic category-specific code; blank for a successful example.", "string", "feature", "Not a production error-code catalog."),
    ("environment", "Logical environment label for the synthetic event.", "string", "feature", "Always simulation in this dataset."),
    ("service_name", "Expected fictional service associated with this example.", "string", "prediction target", "Confirm with correlated logs before treating as fact."),
    ("server_id", "Fictional logical server hosting the expected service.", "string", "prediction target", "Not a resolvable network host."),
    ("request_pattern", "Synthetic method and mock-only route pattern.", "string", "feature", "Routes under /mock are illustrative, not internal production APIs."),
    ("response_pattern", "Expected synthetic response status and event summary.", "string", "prediction target", "Some incomplete-evidence scenarios intentionally have no status."),
    ("http_status", "Expected simulated HTTP status code.", "integer or blank", "prediction target", "Blank when response is intentionally missing."),
    ("historical_root_cause", "Fictional explanation assigned to the generated scenario.", "string", "evaluation-only", "Not a proven cause for any real incident."),
    ("resolution", "Suggested check for the synthetic scenario.", "string", "evaluation-only", "No automated remediation is executed."),
    ("scenario_id", "Unique group key linking workbook rows and log events.", "string", "identifier / split group", "All records for a scenario use one data_split."),
    ("is_synthetic", "Marks that the row is fictional demo data.", "boolean", "quality metadata", "Must be TRUE for every row."),
    ("data_split", "Scenario-level train, validation, or test allocation.", "string", "evaluation control", "Splits are assigned by scenario; no scenario group crosses splits."),
    ("log_directory", "Relative log directory for a fictional service.", "string", "identifier / source location", "Restricted to APP/service_name under the configured mock root."),
    ("enabled", "Whether this fictional registry row is included in simulation.", "boolean", "configuration metadata", "Does not activate a real service."),
    ("service_dependencies", "Semicolon-separated fictional service dependencies.", "string", "feature / configuration metadata", "Describes this synthetic dataset only."),
    ("description", "Plain-language description of the fictional demo service.", "string", "documentation metadata", "Not a production system description."),
    ("scenario_name", "Unique descriptive name of a generated scenario.", "string", "identifier", "One row represents one incident/scenario group."),
    ("input_error", "Synthetic input error used by a scenario.", "string", "feature", "Fictional wording only."),
    ("expected_service", "Fictional service expected at the relevant event point.", "string", "prediction target", "Confirmed against logs only for these fixtures."),
    ("expected_error_category", "Synthetic category label for evaluation.", "string", "prediction target", "Categories are not production taxonomy claims."),
    ("expected_root_cause", "Expected explanation assigned to the fixture.", "string", "evaluation-only", "Not a proven cause for a real incident."),
    ("expected_request_id", "Deterministic ID shared by correlated log events.", "string or blank", "identifier / evaluation-only", "Blank for scenarios without reliable correlation."),
    ("expected_http_status", "Expected final synthetic HTTP status.", "integer or blank", "prediction target", "Blank for intentionally incomplete traces."),
    ("expected_evidence", "Summary of expected generated event evidence.", "string", "evaluation-only", "Used by fixture validation and demos."),
    ("expected_resolution", "Suggested diagnostic check for a synthetic scenario.", "string", "evaluation-only", "No remediation is executed."),
    ("expected_result_status", "Expected diagnostic evidence status.", "string", "evaluation-only", "Partial for incomplete or uncorrelated traces."),
]


def load_registry() -> list[dict[str, Any]]:
    config = json.loads(REGISTRY_PATH.read_text(encoding="utf-8"))
    services = config["services"]
    for row in services:
        if row["application_name"] not in APP_DOMAINS:
            raise ValueError("Registry has an unsupported fictional application.")
        expected = f"{row['application_name']}/{row['service_name']}"
        if row["log_directory"] != expected:
            raise ValueError("Registry log directories must follow APP/service_name.")
    return services


def make_scenario_rows(registry: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    service_index = {(row["application_name"], row["service_name"]): row for row in registry}
    training: list[dict[str, Any]] = []
    catalog: list[dict[str, Any]] = []
    for app_number, application in enumerate(APP_DOMAINS):
        app_rows = SCENARIO_TEMPLATES[application]
        for number, template in enumerate(app_rows, start=1):
            name, error, code, category, service, status, mode, root_cause, resolution = template
            registry_row = service_index[(application, service)]
            scenario_id = f"SCN-{application}-{number:03d}"
            request_id = None if mode == "insufficient" else f"SIM-{application}-{number:04d}"
            split = "train" if number <= 16 else "validation" if number <= 20 else "test"
            method = "POST" if category == "assignment_conflict" else "GET"
            route = {"VNSP": "/mock/vdc/{vdc_id}", "ESP": "/mock/entities/{entity_id}", "IPSM": "/mock/customer-ips/{record_id}"}[application]
            if mode == "assignment_conflict":
                route += "/assignment"
            response = f"HTTP {status}" if status is not None else "No response status recorded"
            if code:
                response += f"; error_code={code}"
            evidence = {
                "success": "Gateway request, domain-service response, and successful gateway response share the request ID.",
                "auth_error": "Gateway request, auth-service rejection, and gateway error response share the request ID.",
                "dependency_error": "Gateway and domain/dependency events share the request ID; response reflects dependency failure.",
                "timeout": "Correlated request, timeout event, and timeout response are present.",
                "connection_refused": "Gateway records refusal and response; downstream service receipt is intentionally absent.",
                "incomplete": "Correlated request and service receipt exist, but response evidence is intentionally absent.",
                "insufficient": "An uncorrelated request event exists without a reliable ID or response evidence.",
                "service_error": "Gateway request, domain-service error/response, and gateway response share the request ID.",
                "gateway_error": "Gateway records the request and synthetic validation or upstream response.",
            }[mode]
            training.append({
                "record_id": f"TRN-{application}-{number:03d}",
                "application_name": application,
                "application_domain": APP_DOMAINS[application],
                "error_message": error,
                "error_code": code,
                "environment": "simulation",
                "service_name": service,
                "server_id": registry_row["server_id"],
                "request_pattern": f"{method} {route}",
                "response_pattern": response,
                "http_status": status,
                "historical_root_cause": f"Fictional demo hypothesis: {root_cause}",
                "resolution": resolution,
                "scenario_id": scenario_id,
                "is_synthetic": True,
                "data_split": split,
            })
            catalog.append({
                "scenario_id": scenario_id,
                "application_name": application,
                "scenario_name": name,
                "input_error": error,
                "expected_service": service,
                "expected_error_category": category,
                "expected_root_cause": root_cause,
                "expected_request_id": request_id,
                "expected_http_status": status,
                "expected_evidence": evidence,
                "expected_resolution": resolution,
                "expected_result_status": "partial_success" if mode in {"incomplete", "insufficient"} else "completed",
                "data_split": split,
                "request_pattern": f"{method} {route}",
                "error_code": code,
                "mode": mode,
            })
    return training, catalog


def _event(
    *, timestamp: datetime, app: str, scenario: dict[str, Any], service: str,
    event_type: str, event_name: str, level: str, message: str,
    registry_index: dict[tuple[str, str], dict[str, Any]],
    http_status: int | None = None, error_code: str | None = None,
    include_id: bool = True,
) -> dict[str, Any]:
    return {
        "timestamp": timestamp.astimezone(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z"),
        "application_name": app,
        "environment": "simulation",
        "service_name": service,
        "server_id": registry_index[(app, service)]["server_id"],
        "level": level,
        "request_id": scenario["expected_request_id"] if include_id else None,
        "correlation_id": scenario["expected_request_id"] if include_id else None,
        "trace_id": f"TRACE-{app}-{scenario['scenario_id'][-3:]}" if include_id else None,
        "scenario_id": scenario["scenario_id"],
        "event_type": event_type,
        "event_name": event_name,
        "message": message,
        "http_method": scenario["request_pattern"].split(" ", 1)[0] if event_type == "request" or service == "api-gateway" else None,
        "endpoint": scenario["request_pattern"].split(" ", 1)[1] if event_type == "request" or service == "api-gateway" else None,
        "http_status": http_status,
        "error_code": error_code or None,
        "is_synthetic": True,
        "data_split": scenario["data_split"],
    }


def make_log_events(catalog: list[dict[str, Any]], registry: list[dict[str, Any]]) -> list[dict[str, Any]]:
    registry_index = {(row["application_name"], row["service_name"]): row for row in registry}
    events: list[dict[str, Any]] = []
    base = datetime(2025, 1, 1, tzinfo=timezone.utc)
    for index, scenario in enumerate(catalog):
        app = scenario["application_name"]
        mode = scenario["mode"]
        domain_service = scenario["expected_service"]
        request_id_missing = mode == "insufficient"
        timestamp = base + timedelta(seconds=index * 30)
        error = scenario["input_error"]
        request_id = scenario["expected_request_id"]
        gateway_message = f"Synthetic request received for {app} detail lookup."
        events.append(_event(
            timestamp=timestamp, app=app, scenario=scenario, service="api-gateway",
            event_type="request", event_name="request_received", level="INFO",
            message=gateway_message, registry_index=registry_index, include_id=not request_id_missing,
        ))
        if mode == "auth_error":
            events.append(_event(
                timestamp=timestamp + timedelta(milliseconds=12), app=app, scenario=scenario,
                service="auth-service", event_type="error", event_name="authorization_rejected",
                level="ERROR", message=error, registry_index=registry_index, error_code=scenario["error_code"], include_id=True,
            ))
        elif mode == "connection_refused":
            events.append(_event(
                timestamp=timestamp + timedelta(milliseconds=18), app=app, scenario=scenario,
                service="api-gateway", event_type="error", event_name="downstream_connection_refused",
                level="ERROR", message=error, registry_index=registry_index, error_code=scenario["error_code"], include_id=True,
            ))
        else:
            domain_level = "ERROR" if mode in {"service_error", "timeout", "incomplete"} else "INFO"
            domain_event = "request_processing" if mode not in {"service_error", "timeout"} else "request_failed"
            events.append(_event(
                timestamp=timestamp + timedelta(milliseconds=20), app=app, scenario=scenario,
                service=domain_service, event_type="request", event_name=domain_event,
                level=domain_level,
                message=(f"Synthetic {app} detail processing started." if mode in {"success", "incomplete", "insufficient"} else error),
                registry_index=registry_index, error_code=scenario["error_code"], include_id=not request_id_missing,
            ))
            if mode == "dependency_error":
                events.append(_event(
                    timestamp=timestamp + timedelta(milliseconds=40), app=app, scenario=scenario,
                    service="mock-dependency", event_type="error", event_name="dependency_unavailable",
                    level="ERROR", message=error, registry_index=registry_index, error_code=scenario["error_code"], include_id=True,
                ))
        if mode not in {"incomplete", "insufficient"}:
            status = scenario["expected_http_status"]
            response_service = "api-gateway"
            response_level = "INFO" if status is not None and status < 400 else "ERROR"
            if mode == "auth_error":
                response_service = "api-gateway"
            elif mode in {"service_error", "timeout", "success", "dependency_error"}:
                response_service = domain_service
            response_time = timestamp + timedelta(milliseconds=80)
            events.append(_event(
                timestamp=response_time, app=app, scenario=scenario, service=response_service,
                event_type="response", event_name="response_sent", level=response_level,
                message=("Synthetic request completed successfully." if status is not None and status < 400 else error),
                registry_index=registry_index, http_status=status,
                error_code=scenario["error_code"], include_id=not request_id_missing,
            ))
            if response_service != "api-gateway":
                events.append(_event(
                    timestamp=response_time + timedelta(milliseconds=10), app=app, scenario=scenario,
                    service="api-gateway", event_type="response", event_name="response_sent",
                    level=response_level, message=("Synthetic gateway returned the service response." if status and status < 400 else error),
                    registry_index=registry_index, http_status=status, error_code=scenario["error_code"], include_id=True,
                ))
    return events


def write_workbook(output: Path, training: list[dict[str, Any]], registry: list[dict[str, Any]], catalog: list[dict[str, Any]]) -> None:
    workbook = Workbook()
    training_sheet = workbook.active
    training_sheet.title = "TrainingData"
    registry_sheet = workbook.create_sheet("ServiceRegistry")
    catalog_sheet = workbook.create_sheet("ScenarioCatalog")
    dictionary_sheet = workbook.create_sheet("DataDictionary")

    tables: list[tuple[Any, str, list[str], list[dict[str, Any]]]] = [
        (training_sheet, "TrainingDataTable", TRAINING_COLUMNS, training),
        (registry_sheet, "ServiceRegistryTable", REGISTRY_COLUMNS, registry),
        (catalog_sheet, "ScenarioCatalogTable", SCENARIO_COLUMNS, catalog),
        (dictionary_sheet, "DataDictionaryTable", DICTIONARY_COLUMNS, [dict(zip(DICTIONARY_COLUMNS, row)) for row in DICTIONARY]),
    ]
    for sheet, table_name, columns, rows in tables:
        sheet.append(columns)
        for row in rows:
            sheet.append([row.get(column) for column in columns])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for column_cells in sheet.columns:
            column_letter = column_cells[0].column_letter
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 14), 58)
        if rows:
            table = Table(displayName=table_name, ref=f"A1:{sheet.cell(row=sheet.max_row, column=len(columns)).coordinate}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            sheet.add_table(table)
        sheet.sheet_view.showGridLines = False
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def generate_dataset(output_root: Path = DEFAULT_OUTPUT) -> dict[str, Any]:
    """Regenerate the workbook, CSV-free JSONL fixture tree, and pending manifest."""
    output_root = Path(output_root).resolve()
    registry = load_registry()
    training, catalog = make_scenario_rows(registry)
    events = make_log_events(catalog, registry)
    logs_root = output_root / "mock_logs"
    by_path: dict[str, list[dict[str, Any]]] = defaultdict(list)
    registry_index = {(row["application_name"], row["service_name"]): row for row in registry}
    for event in events:
        service = registry_index[(event["application_name"], event["service_name"])]
        relative_log_path = f"{service['log_directory']}/app.log"
        by_path[relative_log_path].append(event)

    file_manifest: list[dict[str, Any]] = []
    for relative_path, records in sorted(by_path.items()):
        records.sort(key=lambda row: (row["timestamp"], row["scenario_id"], row["event_name"]))
        local_path = logs_root.joinpath(*Path(relative_path).parts)
        local_path.parent.mkdir(parents=True, exist_ok=True)
        content = "".join(json.dumps(row, sort_keys=True, separators=(",", ":")) + "\n" for row in records)
        local_path.write_text(content, encoding="utf-8", newline="\n")
        file_manifest.append({
            "relative_path": relative_path.replace("\\", "/"),
            "local_path": str(local_path.relative_to(output_root)).replace("\\", "/"),
            "sha256": hashlib.sha256(content.encode("utf-8")).hexdigest(),
            "event_count": len(records),
            "scenario_ids": sorted({row["scenario_id"] for row in records}),
            "upload_status": "not_attempted",
            "upload_message": "Generation only; publishing has not been attempted.",
        })

    workbook_path = output_root / "loglens_training_data.xlsx"
    write_workbook(workbook_path, training, registry, catalog)
    manifest = {
        "manifest_version": 1,
        "data_label": "fictional synthetic demo data; not production incidents",
        "scenario_count": len(catalog),
        "scenario_count_by_application": {app: sum(row["application_name"] == app for row in catalog) for app in APP_DOMAINS},
        "event_count": len(events),
        "workbook": str(workbook_path.relative_to(output_root)).replace("\\", "/"),
        "files": file_manifest,
    }
    (output_root / "upload_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return manifest
