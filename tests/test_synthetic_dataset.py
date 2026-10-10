"""Tests for deterministic workbook/log generation and safe SFTP publication."""

from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from log_providers import LogSourceError, SftpLogProvider
from scripts.mock_dataset import APP_DOMAINS, generate_dataset
from scripts.publish_mock_logs import publish_manifest
from scripts.validate_mock_data import validate_dataset


def test_generates_valid_three_application_workbook_logs_and_grouped_splits(tmp_path):
    output = tmp_path / "training"
    manifest = generate_dataset(output)
    assert manifest["scenario_count"] == 72
    assert manifest["scenario_count_by_application"] == {"VNSP": 24, "ESP": 24, "IPSM": 24}
    assert manifest["event_count"] > manifest["scenario_count"]
    assert len(manifest["files"]) == 12
    assert validate_dataset(output) == []

    workbook = load_workbook(output / "loglens_training_data.xlsx", read_only=True, data_only=True)
    assert workbook.sheetnames == ["TrainingData", "ServiceRegistry", "ScenarioCatalog", "DataDictionary"]
    training_rows = list(workbook["TrainingData"].iter_rows(values_only=True))
    scenario_rows = list(workbook["ScenarioCatalog"].iter_rows(values_only=True))
    training_header = list(training_rows[0])
    scenario_header = list(scenario_rows[0])
    split_col = training_header.index("data_split")
    scenario_id_col = training_header.index("scenario_id")
    synthetic_col = training_header.index("is_synthetic")
    split_by_scenario: dict[str, set[str]] = {}
    for row in training_rows[1:]:
        split_by_scenario.setdefault(row[scenario_id_col], set()).add(row[split_col])
        assert row[synthetic_col] is True
    assert all(len(splits) == 1 for splits in split_by_scenario.values())
    assert sum("test" in splits for splits in split_by_scenario.values()) == 12
    assert scenario_header[0] == "scenario_id"
    assert set(APP_DOMAINS) == {row[scenario_header.index("application_name")] for row in scenario_rows[1:]}
    workbook.close()


def test_generation_is_idempotent_and_contains_success_for_each_domain(tmp_path):
    output = tmp_path / "training"
    first = generate_dataset(output)
    first_checksums = {entry["relative_path"]: entry["sha256"] for entry in first["files"]}
    second = generate_dataset(output)
    second_checksums = {entry["relative_path"]: entry["sha256"] for entry in second["files"]}
    assert first_checksums == second_checksums

    workbook = load_workbook(output / "loglens_training_data.xlsx", read_only=True, data_only=True)
    sheet = workbook["ScenarioCatalog"]
    rows = list(sheet.iter_rows(values_only=True))
    headers = list(rows[0])
    records = [dict(zip(headers, row)) for row in rows[1:]]
    assert all(any(item["application_name"] == app and item["expected_error_category"] == "success" for item in records) for app in APP_DOMAINS)
    assert any(item["expected_error_category"] == "assignment_conflict" for item in records if item["application_name"] == "IPSM")
    workbook.close()


class _MemorySftp:
    def __init__(self):
        self.directories = {"/home/demo"}
        self.files: dict[str, bytes] = {}
        self.closed = False

    def normalize(self, path):
        assert path == "."
        return "/home/demo"

    def stat(self, path):
        if path in self.directories:
            return SimpleNamespace(st_size=0)
        if path in self.files:
            return SimpleNamespace(st_size=len(self.files[path]))
        raise FileNotFoundError(2, "missing")

    def mkdir(self, path):
        parent = str(Path(path).parent).replace("\\", "/")
        if parent not in self.directories:
            raise FileNotFoundError(2, "parent missing")
        self.directories.add(path)

    def open(self, path, mode):
        if mode == "rb":
            if path not in self.files:
                raise FileNotFoundError(2, "missing")
            return io.BytesIO(self.files[path])
        if mode == "xb":
            if path in self.files:
                raise FileExistsError("already exists")
            parent = str(Path(path).parent).replace("\\", "/")
            if parent not in self.directories:
                raise FileNotFoundError(2, "parent missing")
            return _CommitBuffer(self.files, path)
        raise AssertionError(f"Unexpected SFTP open mode: {mode}")

    def rename(self, source, destination):
        if destination in self.files:
            raise FileExistsError("destination exists")
        self.files[destination] = self.files.pop(source)

    def remove(self, path):
        if path not in self.files:
            raise FileNotFoundError(2, "missing")
        del self.files[path]

    def close(self):
        self.closed = True


class _CommitBuffer(io.BytesIO):
    def __init__(self, files, path):
        super().__init__()
        self.files = files
        self.path = path

    def close(self):
        if not self.closed:
            self.files[self.path] = self.getvalue()
        super().close()


def _publisher(tmp_path: Path, sftp: _MemorySftp) -> SftpLogProvider:
    key_file = tmp_path / "demo-key"
    known_hosts = tmp_path / "known_hosts"
    key_file.write_text("fictional placeholder", encoding="utf-8")
    known_hosts.write_text("fictional trusted-host entry", encoding="utf-8")
    provider = SftpLogProvider(
        __import__("infrastructure").InfrastructureManager(),
        host="127.0.0.1", port=2222, username="demo-user", key_filename=key_file,
        known_hosts=known_hosts, remote_root="~/loglens/mock_logs",
        allowed_services={"api-gateway"}, generated_log_root=tmp_path / "training" / "mock_logs",
    )
    provider._connect = lambda: (SimpleNamespace(close=lambda: None), sftp)
    return provider


def test_sftp_publication_is_verified_idempotent_and_refuses_overwrite(tmp_path):
    output = tmp_path / "training"
    manifest = generate_dataset(output)
    entry = next(item for item in manifest["files"] if item["relative_path"] == "VNSP/api-gateway/app.log")
    sftp = _MemorySftp()
    provider = _publisher(tmp_path, sftp)

    first = provider.upload_generated_logs([entry])
    destination = "/home/demo/loglens/mock_logs/VNSP/api-gateway/app.log"
    assert first[0].status == "uploaded"
    assert hashlib.sha256(sftp.files[destination]).hexdigest() == entry["sha256"]
    assert first[0].sha256 == entry["sha256"]

    second = provider.upload_generated_logs([entry])
    assert second[0].status == "unchanged"
    original = sftp.files[destination]
    sftp.files[destination] = b"unrelated remote content"
    conflict = provider.upload_generated_logs([entry])
    assert conflict[0].status == "failed"
    assert "refusing to overwrite" in conflict[0].message
    assert sftp.files[destination] == b"unrelated remote content"
    assert original != sftp.files[destination]


def test_sftp_publisher_rejects_unregistered_or_tampered_manifest_entries(tmp_path):
    output = tmp_path / "training"
    manifest = generate_dataset(output)
    entry = next(item for item in manifest["files"] if item["relative_path"] == "VNSP/api-gateway/app.log")
    sftp = _MemorySftp()
    provider = _publisher(tmp_path, sftp)
    with pytest.raises(LogSourceError, match="outside the application/service publish allowlist"):
        provider.upload_generated_logs([{**entry, "relative_path": "VNSP/../unrelated/app.log"}])
    with pytest.raises(LogSourceError, match="checksum does not match"):
        provider.upload_generated_logs([{**entry, "sha256": "0" * 64}])
    assert not sftp.files


def test_publish_failure_is_written_to_manifest_without_claiming_upload(tmp_path):
    output = tmp_path / "training"
    generate_dataset(output)
    path = output / "upload_manifest.json"

    def disabled_provider():
        raise LogSourceError("SSH log retrieval is disabled or not marked authorized.")

    manifest, ok = publish_manifest(path, provider_factory=disabled_provider)
    assert not ok
    assert manifest["publish_status"] == "failed"
    assert all(item["upload_status"] == "failed" for item in manifest["files"])
    assert all("disabled" in item["upload_message"] for item in manifest["files"])
